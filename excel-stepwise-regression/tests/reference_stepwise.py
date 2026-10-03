"""Line-for-line Python port of StepwiseRegression.bas, used to verify the
algorithm against statsmodels and against hand-run regressions.

Usage: python reference_stepwise.py data.xlsx "Sheet" Y_COL X_COLS FIRST_ROW LAST_ROW
Example (Hoyda workbook): python reference_stepwise.py Hoyda.xlsx "Reg - 5 preds" A B:F 4 48
"""
import math
import sys

from scipy import stats

SING_TOL = 1e-10
MAX_STEPS = 1000


class DataSet:
    def __init__(self, y, X, names):
        self.n, self.k = len(y), len(names)
        self.y, self.x, self.names = y, X, names
        self.ybar = sum(y) / self.n
        self.xbar = [sum(r[j] for r in X) / self.n for j in range(self.k)]
        dy = [v - self.ybar for v in y]
        dx = [[r[j] - self.xbar[j] for j in range(self.k)] for r in X]
        self.syy = sum(v * v for v in dy)
        self.cxy = [sum(dx[i][j] * dy[i] for i in range(self.n)) for j in range(self.k)]
        self.cxx = [[sum(dx[i][a] * dx[i][b] for i in range(self.n)) for b in range(self.k)]
                    for a in range(self.k)]


def invert_spd(a):
    m = len(a)
    a = [row[:] for row in a]
    inv = [[1.0 if i == j else 0.0 for j in range(m)] for i in range(m)]
    d = [a[i][i] for i in range(m)]
    for c in range(m):
        piv = a[c][c]
        if d[c] <= 0 or piv <= SING_TOL * d[c]:
            return None
        for j in range(m):
            a[c][j] /= piv
            inv[c][j] /= piv
        for r in range(m):
            if r != c and a[r][c] != 0:
                f = a[r][c]
                for j in range(m):
                    a[r][j] -= f * a[c][j]
                    inv[r][j] -= f * inv[c][j]
    return inv


def fit(ds, vars_):
    p = len(vars_)
    dfe = ds.n - p - 1
    if dfe < 1:
        return None
    if p == 0:
        return dict(vars=[], coef=[ds.ybar], se=[math.sqrt(ds.syy / dfe / ds.n)], sse=ds.syy, dfe=dfe)
    inv = invert_spd([[ds.cxx[a][b] for b in vars_] for a in vars_])
    if inv is None:
        return None
    b = [sum(inv[i][j] * ds.cxy[vars_[j]] for j in range(p)) for i in range(p)]
    b0 = ds.ybar - sum(b[i] * ds.xbar[vars_[i]] for i in range(p))
    sse = 0.0
    for i in range(ds.n):
        e = ds.y[i] - b0 - sum(b[j] * ds.x[i][vars_[j]] for j in range(p))
        sse += e * e
    mse = sse / dfe
    v0 = 1 / ds.n + sum(ds.xbar[vars_[i]] * inv[i][j] * ds.xbar[vars_[j]] for i in range(p) for j in range(p))
    se = [math.sqrt(mse * v0)] + [math.sqrt(mse * inv[i][i]) for i in range(p)]
    return dict(vars=list(vars_), coef=[b0] + b, se=se, sse=sse, dfe=dfe)


def tstat(c, s):
    if s > 0:
        return c / s
    return 0.0 if c == 0 else math.copysign(1e300, c)


def pval(c, s, df):
    return 2 * stats.t.sf(abs(tstat(c, s)), df)


def current(ds, inm):
    return fit(ds, [j for j in range(ds.k) if inm[j]])


def run(ds, method, p_enter=0.05, p_remove=0.10):
    inm = [method == 2] * ds.k
    log = [("Start", "", None, current(ds, inm))]

    def try_enter():
        best, best_t, best_fit = None, -1, None
        for j in range(ds.k):
            if not inm[j]:
                inm[j] = True
                f = current(ds, inm)
                inm[j] = False
                if f:
                    pos = f["vars"].index(j) + 1
                    t = abs(tstat(f["coef"][pos], f["se"][pos]))
                    if t > best_t:
                        best, best_t, best_fit = j, t, f
        if best is None:
            return False
        pos = best_fit["vars"].index(best) + 1
        pv = pval(best_fit["coef"][pos], best_fit["se"][pos], best_fit["dfe"])
        if pv < p_enter:
            inm[best] = True
            log.append(("Entered", ds.names[best], pv, best_fit))
            return True
        return False

    def try_remove():
        f = current(ds, inm)
        if not f or not f["vars"]:
            return False
        ts = [abs(tstat(f["coef"][i], f["se"][i])) for i in range(1, len(f["vars"]) + 1)]
        w = ts.index(min(ts)) + 1
        pv = pval(f["coef"][w], f["se"][w], f["dfe"])
        if pv > p_remove:
            inm[f["vars"][w - 1]] = False
            log.append(("Removed", ds.names[f["vars"][w - 1]], pv, current(ds, inm)))
            return True
        return False

    guard = 0
    while True:
        changed = False
        if method != 2 and try_enter():
            changed = True
        if method != 1:
            while try_remove():
                changed = True
                guard += 1
                if guard > MAX_STEPS:
                    break
        guard += 1
        if not (changed and guard <= MAX_STEPS):
            break
    return inm, log


def describe(ds, inm, log):
    for i, (act, name, pv, f) in enumerate(log):
        r2 = 1 - f["sse"] / ds.syy
        print(f"  step {i}: {act:8s} {name:12s} p={'' if pv is None else f'{pv:.6g}':12s} "
              f"terms={len(f['vars'])} R2={r2:.6f}")
    f = current(ds, inm)
    print("  final:", ["Intercept"] + [ds.names[j] for j in f["vars"]])
    for i, c in enumerate(f["coef"]):
        print(f"    coef={c:.10g} se={f['se'][i]:.10g} p={pval(c, f['se'][i], f['dfe']):.6g}")


def load(path, sheet, ycol, xcols, first, last):
    import openpyxl
    from openpyxl.utils import column_index_from_string as ci
    ws = openpyxl.load_workbook(path, data_only=True)[sheet]
    a, b = (xcols.split(":") + [xcols])[:2]
    xs = list(range(ci(a), ci(b) + 1))
    names = [ws.cell(first, c).value for c in xs]
    y, X = [], []
    for r in range(first + 1, last + 1):
        yv = ws.cell(r, ci(ycol)).value
        row = [ws.cell(r, c).value for c in xs]
        if all(isinstance(v, (int, float)) and not isinstance(v, bool) for v in [yv] + row):
            y.append(float(yv))
            X.append([float(v) for v in row])
    return DataSet(y, X, names)


if __name__ == "__main__":
    ds = load(sys.argv[1], sys.argv[2], sys.argv[3], sys.argv[4], int(sys.argv[5]), int(sys.argv[6]))
    for m, name in [(1, "Forward"), (2, "Backward"), (3, "Stepwise")]:
        print(name)
        inm, log = run(ds, m)
        describe(ds, inm, log)
