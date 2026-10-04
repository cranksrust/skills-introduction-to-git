"""Writes examples/Stepwise_Example_Output.xlsx: the add-in's report, styling
included, for forward, backward and stepwise runs (t Stat criterion) on the
Hoyda data. Mirrors WriteReport in StepwiseRegression.bas.

Usage: python make_example.py path/to/Hoyda_Discrimination.xlsx output.xlsx
"""
import math
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from scipy import stats

import reference_stepwise as ref

ORANGE = "BF5700"
TINT = "F8E5D6"
GREY = "595959"
CREDIT = "Built for the HOUMBA Class of '28"
FMT_P = '[<0.0001]"<0.0001";0.0000'
FMT_T, FMT_R2, FMT_NUM, FMT_COEF = "0.00", "0.0000", "#,##0.00", "#,##0.0000"


def font(**kw):
    return Font(name="Calibri", size=kw.pop("size", 10), **kw)


def write_report(wb, ds, method, mname, sname):
    inm, log = ref.run(ds, method, ref.CRIT_T)
    ws = wb.create_sheet(sname)
    ws.sheet_view.showGridLines = False
    r = [4]

    def put(row, col, v, fmt=None, **kw):
        c = ws.cell(row, col, v)
        c.font = font(**kw)
        if fmt:
            c.number_format = fmt
        return c

    def section(title):
        put(r[0], 1, title, bold=True, size=11, color=ORANGE)
        r[0] += 1

    def pair(label, v, fmt=None):
        put(r[0], 1, label, color=GREY)
        put(r[0], 2, v, fmt).alignment = Alignment(horizontal="left")
        r[0] += 1

    def header(hs, fmts):
        for i, (h, f) in enumerate(zip(hs, fmts)):
            c = put(r[0], i + 1, h, bold=True)
            c.fill = PatternFill("solid", fgColor=TINT)
            c.border = Border(bottom=Side(style="medium", color=ORANGE))
            c.alignment = Alignment(horizontal="left" if f == "@" else "right", vertical="bottom", wrap_text=True)
        r[0] += 1

    def row(vals, fmts):
        for c, (v, f) in enumerate(zip(vals, fmts)):
            if v is not None:
                cell = put(r[0], c + 1, v, None if f == "@" else f)
                if f == "@":
                    cell.alignment = Alignment(horizontal="left")
        r[0] += 1

    for c in range(1, 12):
        cell = ws.cell(1, c)
        cell.fill = PatternFill("solid", fgColor=ORANGE)
    put(1, 1, "  Stepwise Regression Output", bold=True, size=14, color="FFFFFF").alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 26
    put(2, 1, "  " + CREDIT, italic=True, size=9, color=ORANGE)

    section("Settings")
    pair("Dependent variable", "SALARY")
    pair("Method", mname)
    pair("Criterion", "t Stat")
    if method != 2:
        pair("|t| to enter", 2, "0.00")
    if method != 1:
        pair("|t| to remove", 2, "0.00")
    pair("Observations used", ds.n, "0")
    pair("Rows excluded (blank or non-numeric)", 0, "0")
    pair("Y range", "'[Hoyda Discrimination.xlsx]Reg - 5 preds'!$A$4:$A$48")
    pair("X range", "'[Hoyda Discrimination.xlsx]Reg - 5 preds'!$B$4:$F$48")
    r[0] += 1

    section("Selection steps")
    fm = ["@", "@", "@", FMT_T, FMT_P, "0", FMT_R2, FMT_R2, FMT_NUM, FMT_NUM, FMT_NUM]
    header(["Step", "Action", "Variable", "t Stat", "P-value", "Terms in model", "R Square",
            "Adjusted R Square", "RMSE", "AIC", "BIC"], fm)
    for i, (act, name, tv, pv, f) in enumerate(log):
        p, sse = len(f["vars"]), f["sse"]
        row([i, act, name or ("(all candidates)" if method == 2 else "(intercept only)"), tv, pv, p,
             1 - sse / ds.syy, 1 - (sse / f["dfe"]) / (ds.syy / (ds.n - 1)), math.sqrt(sse / f["dfe"]),
             ds.n * math.log(sse / ds.n) + 2 * (p + 1), ds.n * math.log(sse / ds.n) + (p + 1) * math.log(ds.n)], fm)

    f = ref.current(ds, inm)
    p, sse, dfe = len(f["vars"]), f["sse"], f["dfe"]
    ssr = ds.syy - sse
    r[0] += 1
    section("Final model")
    pair("Multiple R", math.sqrt(1 - sse / ds.syy), FMT_R2)
    pair("R Square", 1 - sse / ds.syy, FMT_R2)
    pair("Adjusted R Square", 1 - (sse / dfe) / (ds.syy / (ds.n - 1)), FMT_R2)
    pair("RMSE", math.sqrt(sse / dfe), FMT_NUM)
    pair("Observations", ds.n, "0")
    r[0] += 1

    fm = ["@", "0", FMT_NUM, FMT_NUM, FMT_NUM, FMT_P]
    header(["ANOVA", "df", "SS", "MS", "F", "Significance F"], fm)
    fs = (ssr / p) / (sse / dfe)
    row(["Regression", p, ssr, ssr / p, fs, stats.f.sf(fs, p, dfe)], fm)
    row(["Residual", dfe, sse, sse / dfe, None, None], fm)
    for c in range(1, 7):
        ws.cell(r[0], c).border = Border(top=Side(style="thin", color="BFBFBF"))
    row(["Total", ds.n - 1, ds.syy, None, None, None], fm)
    r[0] += 1

    fm = ["@", FMT_COEF, FMT_COEF, FMT_T, FMT_P, FMT_COEF, FMT_COEF]
    header(["Term", "Coefficients", "Standard Error", "t Stat", "P-value", "Lower 95%", "Upper 95%"], fm)
    tc = stats.t.ppf(0.975, dfe)
    for i, nm in enumerate(["Intercept"] + [ds.names[j] for j in f["vars"]]):
        b, s = f["coef"][i], f["se"][i]
        row([nm, b, s, b / s, ref.pval(b, s, dfe), b - tc * s, b + tc * s], fm)
    r[0] += 1

    section("Variables not in the final model")
    fm = ["@", FMT_T, FMT_P, "@"]
    header(["Variable", "t Stat if added", "P-value if added", "Note"], fm)
    for j in range(ds.k):
        if not inm[j]:
            inm[j] = True
            t = ref.current(ds, inm)
            inm[j] = False
            pos = t["vars"].index(j) + 1
            row([ds.names[j], ref.tstat(t["coef"][pos], t["se"][pos]),
                 ref.pval(t["coef"][pos], t["se"][pos], t["dfe"]), None], fm)

    r[0] += 1
    for c in range(1, 12):
        ws.cell(r[0], c).border = Border(top=Side(style="thin", color=ORANGE))
    put(r[0], 1, "Stepwise Regression add-in  |  " + CREDIT, italic=True, size=8, color=GREY)

    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 13
    ws.column_dimensions["C"].width = 16
    for col in "DEFGHIJK":
        ws.column_dimensions[col].width = 12


if __name__ == "__main__":
    ds = ref.load(sys.argv[1], "Reg - 5 preds", "A", "B:F", 4, 48)
    wb = Workbook()
    wb.remove(wb.active)
    write_report(wb, ds, 1, "Forward selection", "Stepwise")
    write_report(wb, ds, 2, "Backward elimination", "Stepwise (2)")
    write_report(wb, ds, 3, "Stepwise", "Stepwise (3)")
    wb.save(sys.argv[2])
